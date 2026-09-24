#!/usr/bin/env python3
"""Build clearly marked method drafts, optionally adding verified saved results.

Use the Codex bundled Python runtime (python-docx required). Presentation output
uses the bundled @oai/artifact-tool JavaScript runtime, never python-pptx.
The builder never starts training or test evaluation. Synthetic runs are refused.
"""

from __future__ import annotations

import argparse
import csv
from datetime import date
import hashlib
import json
import math
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile

ROOT = Path(__file__).resolve().parents[1]
DEPENDENCIES = Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies"
SKILLS = Path.home() / ".codex/plugins/cache/openai-primary-runtime"
FAMILIES = {"svm": "LinearSVC", "logistic_regression": "Logistic Regression", "phobert": "PhoBERT", "bamibert": "BamiBERT", "complement_nb": "ComplementNB (phụ)"}
SOURCES = [
    ("ViHSD của UIT", "https://huggingface.co/datasets/uitnlp/vihsd"),
    ("Repository tác giả ViHSD", "https://github.com/sonlam1102/vihsd"),
    ("Hiệu chỉnh xác suất", "https://scikit-learn.org/stable/modules/calibration.html"),
    ("Lựa chọn ngưỡng", "https://scikit-learn.org/stable/modules/classification_threshold.html"),
    ("So sánh mô hình văn bản", "https://scikit-learn.org/stable/auto_examples/text/plot_document_classification_20newsgroups.html"),
    ("PhoBERT checkpoint gốc", "https://huggingface.co/vinai/phobert-base"),
    ("BamiBERT checkpoint gốc", "https://huggingface.co/Qualcomm-AI-Research/BamiBERT"),
]


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def saved_results(run_dir: Path | None, final_dir: Path | None) -> dict | None:
    if final_dir and not run_dir:
        raise ValueError("--final-dir requires --run-dir so the locked selection can be verified.")
    if not run_dir:
        return None
    metadata = read_json(run_dir / "metadata.json")
    selection = read_json(run_dir / "selection.json")
    if metadata.get("synthetic") is not False or selection.get("synthetic") is not False or metadata.get("dataset_manifest", {}).get("synthetic", False):
        raise ValueError("Reports refuse synthetic or ambiguously marked runs. These are not ViHSD findings.")
    if metadata.get("status") != "locked_awaiting_test":
        raise ValueError("The run must have a completed, locked selection.")
    if metadata.get("dataset_fingerprint") != selection.get("dataset_fingerprint"):
        raise ValueError("Run and selection fingerprints do not match.")
    if metadata.get("run_id") != selection.get("run_id") or selection.get("selected_family") not in FAMILIES:
        raise ValueError("Run identity or selected model is inconsistent.")
    if hashlib.sha256((run_dir / "config.json").read_bytes()).hexdigest() != selection.get("config_sha256"):
        raise ValueError("Saved configuration changed after selection was locked.")
    config = read_json(run_dir / "config.json")
    families = set(config.get("models", {}))
    if not families or not families.issubset(FAMILIES) or selection["selected_family"] not in families:
        raise ValueError("Saved configuration contains missing or unsupported model families.")
    declared = config.get("deployment_family")
    if declared is not None and declared != selection["selected_family"]:
        raise ValueError("Selected deployment family does not match the predeclared configuration.")
    candidates = selection.get("candidates", [])
    if len(candidates) != len(families) or {item["family"] for item in candidates} != families:
        raise ValueError("Frozen candidates do not match the configured model families.")
    with (run_dir / "comparison.csv").open(encoding="utf-8", newline="") as handle:
        comparison = list(csv.DictReader(handle))
    if len(comparison) != len(families) or {row["family"] for row in comparison} != families:
        raise ValueError("Expected exactly one comparison row for each configured model family.")
    for row in comparison:
        score = float(row["validation_macro_f1"])
        if not math.isfinite(score) or not 0 <= score <= 1:
            raise ValueError("Invalid saved validation Macro-F1.")
    best_family = min(comparison, key=lambda row: (-float(row["validation_macro_f1"]), row["family"]))["family"]
    if selection.get("validation_best_family", selection["selected_family"]) != best_family:
        raise ValueError("Saved validation winner does not match the comparison table.")
    result = {"run_dir": str(run_dir), "metadata": metadata, "selection": selection, "validation": comparison}
    if final_dir:
        final_metadata = read_json(final_dir / "metadata.json")
        expected_hash = hashlib.sha256((run_dir / "selection.json").read_bytes()).hexdigest()
        if final_metadata.get("synthetic") is not False or final_metadata.get("status") != "evaluated":
            raise ValueError("Final evaluation must be completed and explicitly non-synthetic.")
        if final_metadata.get("selection_sha256") != expected_hash or final_metadata.get("dataset_fingerprint") != selection["dataset_fingerprint"]:
            raise ValueError("Final evaluation does not match the frozen selection and dataset.")
        with (final_dir / "comparison.csv").open(encoding="utf-8", newline="") as handle:
            result["test"] = list(csv.DictReader(handle))
        if len(result["test"]) != len(families) or {row["family"] for row in result["test"]} != families:
            raise ValueError("Expected exactly one test row for each frozen model family.")
        for row in result["test"]:
            for metric in ("macro_f1", "accuracy", "weighted_f1"):
                value = float(row[metric])
                if not math.isfinite(value) or not 0 <= value <= 1:
                    raise ValueError("Invalid saved test metric.")
        result["final_dir"] = str(final_dir)
    return result


def representation_label(row):
    if row["family"] == "phobert":
        return "PyVi + tokenizer"
    if row["family"] == "bamibert":
        return "Tokenizer (raw)"
    return f"TF-IDF {row['feature_kind']}"


def selection_description(selection):
    selected = FAMILIES[selection["selected_family"]]
    best = FAMILIES[selection.get("validation_best_family", selection["selected_family"])]
    reason = "đã khai báo trước" if selection.get("selection_criterion") == "predeclared_deployment_family" else "chọn bằng validation"
    return f"Triển khai: {selected} ({reason}). Validation cao nhất: {best}."


def set_cell_shading(cell, color):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(node)


def add_table(document, headings, rows, widths):
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    table = document.add_table(rows=1, cols=len(headings))
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    for cell, width in zip(table.rows[0].cells, widths):
        cell.width = Cm(width)
    for index, label in enumerate(headings):
        table.cell(0, index).text = label
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].width = Cm(widths[index])
            cells[index].text = str(value)
    props = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            border.set(qn(f"w:{key}"), value)
        borders.append(border)
    props.append(borders)
    for row_number, row in enumerate(table.rows):
        for column, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins = OxmlElement("w:tcMar")
            for side in ("top", "left", "bottom", "right"):
                margin = OxmlElement(f"w:{side}")
                margin.set(qn("w:w"), "100")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            cell._tc.get_or_add_tcPr().append(margins)
            set_cell_shading(cell, "DCE6EE" if row_number == 0 else ("F6F8FA" if row_number % 2 == 0 else "FFFFFF"))
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.space_before = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                    run.bold = row_number == 0
                    run.font.color.rgb = RGBColor(0, 0, 0)
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    document.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def paragraph(doc, text, *, bold=False):
    p = doc.add_paragraph()
    p.add_run(text).bold = bold
    return p


def build_docx(path, results):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.8), Cm(1.7)
    section.left_margin, section.right_margin = Cm(2), Cm(2)
    section.header_distance, section.footer_distance = Cm(.8), Cm(.8)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Header", "Footer"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(10.5 if name == "Normal" else 10)
    # Remove decorative rules inherited from the bundled Word default template.
    for element in doc.styles.element.xpath(".//w:pBdr"):
        element.getparent().remove(element)
    normal = doc.styles["Normal"].paragraph_format
    normal.space_after, normal.line_spacing = Pt(7), 1.12
    doc.styles["Title"].font.size = Pt(25)
    doc.styles["Title"].font.bold = True
    doc.styles["Heading 1"].font.size = Pt(18)
    doc.styles["Heading 2"].font.size = Pt(12)
    for name in ("Heading 1", "Heading 2"):
        doc.styles[name].font.bold = True
        doc.styles[name].paragraph_format.space_after = Pt(9)
        doc.styles[name].paragraph_format.space_before = Pt(10)
    header = section.header.paragraphs[0]
    header.text = "CS114     SafeView     Bản nháp phương pháp"
    header.runs[0].font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Bản nháp  |  Trang ").font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = "Bản nháp phương pháp phân loại bình luận tiếng Việt cho SafeView"
    doc.core_properties.subject = "CS114; bản nháp phương pháp; chưa phải báo cáo kết quả hoàn chỉnh"
    doc.core_properties.author = "Nhóm đồ án CS114"

    doc.add_paragraph("Bản nháp phương pháp\nphân loại bình luận tiếng Việt\ncho SafeView", style="Title")
    paragraph(doc, "BẢN NHÁP DRAFT", bold=True)
    status = "Chưa có kết quả thực nghiệm ViHSD. Nội dung hiện tại mô tả phương pháp và phần mềm đã triển khai." if results is None else "Có bảng số liệu đọc từ run đã lưu ở phụ lục. Phân tích lỗi và kết luận học thuật vẫn cần người thực hiện rà soát."
    paragraph(doc, status, bold=True)
    paragraph(doc, f"Cập nhật bản tài liệu {date.today().isoformat()}. Dành cho nhóm đồ án và giảng viên khi rà soát phạm vi, thiết kế thực nghiệm và các phần còn thiếu trước khi nộp.")
    doc.add_heading("Mục tiêu nghiên cứu", level=1)
    paragraph(doc, "Thiết kế chính so sánh LinearSVC, Logistic Regression, PhoBERT và BamiBERT trên cùng split ViHSD, với đầu ra CLEAN, OFFENSIVE và HATE. Hai baseline dùng TF-IDF; hai Transformer được fine-tune từ checkpoint gốc. ComplementNB được giữ làm baseline phụ trong cấu hình legacy khi cần ba thuật toán truyền thống.")
    paragraph(doc, "Macro-F1 validation dùng để chọn cấu hình từng họ và ghi nhận mô hình đứng đầu. BamiBERT là đích triển khai SafeView đã khai báo trước; không gọi đó là mô hình tốt nhất nếu kết quả không chứng minh. Khóa model và chính sách ẩn trước khi đánh giá test. Chất lượng và độ trễ triển khai cần được đo thực tế.")
    add_table(doc, ["Hạng mục", "Trạng thái tài liệu này"], [
        ["Package và workflow", "Đã có mã nạp dữ liệu, EDA, train, khóa lựa chọn, đánh giá và serving"],
        ["ViHSD và số liệu", "Loader GitHub đã có; chưa chạy thực nghiệm và chưa có điểm số" if results is None else "Đọc từ run đã chỉ định ở phụ lục"],
        ["Kiểm thử phần mềm", "Mẫu tổng hợp chỉ kiểm tra cơ chế, không dùng làm kết quả học thuật"],
        ["Deploy và bài nộp cuối", "Cần model thật, đánh giá và xác nhận vận hành trước khi hoàn tất"],
    ], [4.5, 12.5])
    paragraph(doc, "Phạm vi chỉ gồm phân loại văn bản đơn lẻ. Đồ án chưa đánh giá ảnh, video, tin giả hoặc đầy đủ ngữ cảnh hội thoại. Nhóm cần xác nhận với giảng viên rằng dữ liệu NLP đáp ứng rubric môn học.")

    doc.add_page_break()
    doc.add_heading("Dữ liệu và kiểm tra split", level=1)
    dataset_status = "Chưa có kết quả ViHSD; người thực hiện sẽ tự chạy notebook và đánh giá." if results is None else "Bản dữ liệu của run đã lưu được xác định bằng fingerprint ở phụ lục. Nhóm cần rà manifest và biểu đồ EDA tương ứng khi diễn giải kết quả."
    paragraph(doc, f"ViHSD cung cấp nhãn CLEAN, OFFENSIVE và HATE cho bình luận mạng xã hội. Notebook nạp trực tiếp data/vihsd.zip từ GitHub tác giả vào RAM, không yêu cầu thư mục corpus local hay đăng nhập Hugging Face mặc định. Cache ZIP trong runtime là tùy chọn. {dataset_status} Nguồn [2]; đường tải Hugging Face [1] vẫn là tùy chọn có xác thực.")
    add_table(doc, ["Giá trị", "Nhãn", "Cách sử dụng"], [["0", "CLEAN", "Lớp sạch theo quy tắc gán nhãn của dataset"], ["1", "OFFENSIVE", "Lớp xúc phạm theo quy tắc dataset"], ["2", "HATE", "Lớp thù ghét theo quy tắc dataset"]], [2, 4, 11])
    paragraph(doc, "Loader giải quyết branch/tag thành commit SHA trước khi đọc ZIP, giữ train/dev/test và chỉ đổi tên dev thành validation. Manifest lưu URL nguồn, commit, hash archive/file, fingerprint nội dung, nhãn và số mẫu. Nếu thiếu ID thì sinh ID từ split và vị trí. Local loader vẫn đọc source.json khi có. ID lỗi, nhãn lỗi hoặc thiếu split đều bị từ chối; không tự thay dataset.")
    doc.add_heading("Audit trước khi huấn luyện", level=2)
    paragraph(doc, "Audit đếm null/rỗng, trùng raw, trùng sau NFC/khoảng trắng, nhóm nhãn mâu thuẫn và trùng giữa split. Kiểm tra gần trùng dùng cosine TF-IDF ký tự trên tối đa 1.500 text khác nhau và 2.000 ký tự mỗi text. Đây là sàng lọc có giới hạn, chưa phải kiểm tra đầy đủ corpus.")
    paragraph(doc, "Không âm thầm xóa hoặc chuyển mẫu. Nếu có leakage, báo rõ kết quả split gốc và thí nghiệm loại trùng bằng manifest khác. Văn bản null/rỗng vẫn được audit nhưng chặn bước train cho đến khi nhóm ghi nhận chính sách xử lý.")
    doc.add_heading("Các đầu ra EDA", level=2)
    add_table(doc, ["Phân tích", "Đầu ra cần đọc"], [["Phân bố nhãn", "Số mẫu mỗi nhãn trên từng split"], ["Độ dài", "Phân bố ký tự và token khoảng trắng theo nhãn"], ["Nội dung train", "Tần suất token 1-gram và 2-gram"], ["Dấu hiệu văn bản", "Emoji, URL, ký tự lặp và chuỗi Latin không dấu"]], [5, 12])
    paragraph(doc, "CSV/JSON/PNG công bố chỉ chứa thống kê tổng hợp. N-gram phải xuất hiện trong ít nhất hai tài liệu và không bằng nguyên một bình luận. Corpus ở RAM hoặc cache runtime tùy chọn; dự đoán từng mẫu dùng phân tích riêng và không đưa vào artifact public. Cần đối chiếu điều kiện nguồn trước khi tái phân phối.")

    doc.add_page_break()
    doc.add_heading("Tiền xử lý và bốn mô hình", level=1)
    paragraph(doc, "Pipeline mặc định chuẩn hóa Unicode NFC và khoảng trắng. Giữ dấu, emoji, phủ định và chữ hoa. Viết thường hoặc thay URL/mention là tùy chọn cấu hình được lưu cùng run và dùng lại khi serving. Toàn bộ vectorizer chỉ fit trên train, hoặc trên phần train của từng fold calibration.")
    add_table(doc, ["Biểu diễn", "Thiết lập", "Giới hạn mặc định"], [["Token", "TF-IDF n-gram 1 đến 2", "50.000 đặc trưng"], ["Ký tự", "TF-IDF n-gram 3 đến 5", "50.000 đặc trưng"], ["Kết hợp", "FeatureUnion của hai nhóm", "Tối đa 100.000 đặc trưng"]], [3.5, 7, 6.5])
    paragraph(doc, "Token tách bằng khoảng trắng thường là âm tiết tiếng Việt, chưa phải từ được phân đoạn bằng bộ tách từ. Ma trận luôn giữ dạng sparse. Character n-gram có thể giúp với teencode hoặc sai chính tả, nhưng đây là giả thuyết cần đo, chưa phải kết luận.")
    add_table(doc, ["Mô hình", "Đầu vào", "Thiết lập khởi đầu"], [["LinearSVC", "TF-IDF sparse", "C: 0,1; 1; 5\nweight: None/balanced"], ["Logistic Regression", "TF-IDF sparse", "C: 0,1; 1; 5\nweight: None/balanced"], ["PhoBERT", "PyVi rồi tokenizer gốc", "LR 2e-5; max_length 128"], ["BamiBERT", "Văn bản trực tiếp và tokenizer gốc", "LR 2e-5; max_length 128"]], [4.2, 6.4, 6.4])
    doc.add_heading("Ngân sách và tiền xử lý theo họ", level=2)
    paragraph(doc, "Baseline thử word/char/combined và min_df trong {2, 3}, rồi tune classifier trên biểu diễn tốt nhất. PhoBERT dùng vinai/phobert-base cùng PyVi; PyVi khác VnCoreNLP dùng trong tiền huấn luyện, cần ghi nhận hạn chế này. BamiBERT dùng Qualcomm-AI-Research/BamiBERT, nhận văn bản sau chuẩn hóa chung và không tách từ. Tokenizer, commit checkpoint và preprocessing đi cùng artifact. Nguồn [5], [6], [7].")
    paragraph(doc, "Transformer mặc định 3 epoch, batch 8, gradient accumulation 2, FP16 khi có CUDA; chọn checkpoint theo validation Macro-F1. Đây là cấu hình khởi đầu cho Colab Free, chưa bảo đảm GPU, thời gian hay chất lượng. Ghi thời gian, bộ nhớ, loss/learning rate và metrics từng epoch; baseline ghi số đặc trưng và warnings. Cấu hình thực tế nằm trong run/config.json.")

    doc.add_page_break()
    doc.add_heading("Validation calibration và khóa test", level=1)
    paragraph(doc, "Validation là nguồn duy nhất để chọn cấu hình và ngưỡng. LinearSVC được hiệu chỉnh sigmoid bằng CalibratedClassifierCV với các fold của train, bọc toàn pipeline để mỗi fold tự fit TF-IDF. Mặc định dùng ba fold và ensemble=False. Sau calibration, đo lại validation vì xác suất mới có thể làm đổi nhãn và thứ hạng. Nguồn [3].")
    paragraph(doc, "LR có predict_proba; Transformer dùng softmax trên logits theo thứ tự nhãn cố định. Cần xem reliability diagram, log loss và Brier score, không mặc định xem softmax là đã hiệu chỉnh. ComplementNB phụ cũng cần kiểm tra xác suất nếu sử dụng. Không tự biến decision score của SVM thành xác suất.")
    add_table(doc, ["Giai đoạn", "Dữ liệu và quyết định"], [["Fit và calibration", "Train; calibration SVM dùng các fold của train"], ["Lựa chọn", "Validation chọn cấu hình/ngưỡng; ghi riêng họ tốt nhất"], ["Khóa", "Đích BamiBERT, preprocessing, nhãn, policy và hashes"], ["Đánh giá cuối", "Test các ứng viên đã khóa; không đổi đích triển khai dựa trên test"]], [4.5, 12.5])
    paragraph(doc, "Run lưu selection.json và SHA-256 của config/artifacts. Đánh giá test kiểm tra lại fingerprint và hashes trước khi chạy. Đã có kết quả cuối thì nạp kết quả đã lưu; workflow chặn chạy lại final hoặc tune lại cùng fingerprint trong cùng thư mục kết quả. Quy trình vẫn cần kỷ luật nghiên cứu khi sao chép dự án sang nơi khác.")
    doc.add_heading("Chỉ số và phân tích", level=2)
    paragraph(doc, "Tiêu chí chính là Macro-F1 ba lớp theo argmax. Báo thêm Accuracy, weighted-F1, Precision/Recall/F1 từng lớp và confusion matrix. Baseline có learning curves theo kích thước train; Transformer có đường loss/metrics theo epoch. Phải ghi rõ hai loại đường này khác nhau. Khi báo chênh lệch F1, dùng điểm phần trăm và nói rõ tập đánh giá.")
    paragraph(doc, "Đo latency trên cùng CPU và batch size, bao gồm preprocessing cùng TF-IDF hoặc tokenization/inference tương ứng; báo p95, kích thước artifact và điều kiện đo. Không so số đặc trưng TF-IDF với số tham số Transformer như cùng một đơn vị. Bootstrap test mô tả bất định do lấy mẫu, chưa bao gồm bất định huấn luyện/chọn cấu hình.")
    paragraph(doc, "Dành khoảng 100 lỗi validation cho rà soát thủ công: chửi đùa, trích dẫn, phủ định, mỉa mai, không dấu/teencode, cần ngữ cảnh và ranh giới OFFENSIVE/HATE. Báo số lượng theo nhóm và nhận xét tổng hợp; không đưa raw bình luận vào bản public.")

    doc.add_page_break()
    doc.add_heading("Chính sách ẩn và tích hợp SafeView", level=1)
    paragraph(doc, "Phân loại ba lớp và quyết định ẩn là hai đầu ra riêng. Nhãn báo cáo lấy argmax của CLEAN/OFFENSIVE/HATE. Quyết định ẩn dùng tổng xác suất của hai lớp harmful, không thêm điều kiện argmax phải thuộc harmful.")
    paragraph(doc, "p_harm = P(OFFENSIVE) + P(HATE)", bold=True)
    paragraph(doc, "Ẩn khi p_harm ≥ threshold", bold=True)
    paragraph(doc, "Mặc định chọn threshold tối đa hóa F1 harmful trên validation. Nếu bằng nhau, ưu tiên FPR CLEAN thấp hơn rồi ngưỡng cao hơn. Threshold sweep lưu Precision, Recall và F1 harmful, tỷ lệ CLEAN bị ẩn nhầm, recall ẩn được HATE và tỷ lệ bị ẩn. Mục tiêu FPR có ràng buộc là tùy chọn sản phẩm, chưa mặc định 5%. Nguồn phương pháp [4].")
    paragraph(doc, "decision_policy.json lưu score, threshold, tiêu chí chọn, metrics validation, label mapping và model revision. Thay preprocessing, calibration hoặc retrain train+validation có thể đổi phân bố score; cần một quy trình lựa chọn mới, không tái sử dụng ngưỡng và kết quả cũ như cùng một pipeline.")
    doc.add_heading("Hợp đồng triển khai", level=2)
    paragraph(doc, "HF Space nạp artifact BamiBERT gồm model, tokenizer, preprocessing và policy đã khóa; Model Hub lưu checkpoint. Gradio endpoint classify trả map phẳng CLEAN/OFFENSIVE/HATE để tương thích provider. SafeView gửi một bình luận, đọc kết quả SSE và áp dụng policy đúng revision. Việc chọn BamiBERT để deploy không thay thế bảng xếp hạng validation.")
    add_table(doc, ["Điểm kiểm tra", "Yêu cầu cho demo"], [["Luật ẩn", "Dùng tổng OFFENSIVE + HATE cùng threshold đã chọn"], ["Lexicon và routing", "Đo riêng ảnh hưởng override; có ca tiếng Việt không dấu"], ["API lỗi hoặc chậm", "Báo trạng thái chưa sẵn sàng; không coi lỗi mạng là CLEAN"], ["Đo vận hành", "Warm/cold latency, p95, concurrency và cache"]], [5, 12])
    paragraph(doc, "Không nhúng HF access token trong extension. Model thuần và toàn SafeView cần bảng kết quả riêng vì routing/lexicon có thể đổi quyết định. Bản nháp này chưa xác nhận deployment public hay cải thiện so với PhoBERT. Các bước tích hợp chi tiết nằm trong docs/safeview-integration.md.")

    doc.add_page_break()
    doc.add_heading("Tái lập và công việc còn lại", level=1)
    paragraph(doc, "Repository tổ chức một notebook duy nhất từ EDA đến kết quả cuối. Notebook và serving cùng gọi safeview_ml để tránh sao chép công thức chuẩn hóa hoặc suy luận. Mỗi run có thư mục riêng, lưu config, môi trường, provenance và các bảng/hình dùng chung cho báo cáo.")
    add_table(doc, ["Bước", "Đầu ra cần hoàn thành"], [["1  Nạp dữ liệu", "GitHub ZIP trong RAM; lưu manifest và commit nguồn"], ["2  EDA và train", "Người dùng chạy bốn mô hình; giữ checkpoint để resume"], ["3  Khóa và đánh giá", "Ghi đích triển khai, họ tốt nhất validation, test và phân tích lỗi"], ["4  Demo và bài nộp", "Model card, Space, SafeView và Word/slide hoàn chỉnh"]], [5, 12])
    paragraph(doc, "Các kiểm thử với mẫu tổng hợp xác nhận hợp đồng dữ liệu, độ tái lập, serialization, policy và API. Điểm số của smoke run không phản ánh khả năng tổng quát hóa trên ViHSD. Bài nộp chỉ hoàn tất sau khi chạy dữ liệu thật, lưu output notebook và viết nhận xét dựa trên kết quả.")
    doc.add_heading("Cách bổ sung bảng kết quả thật", level=2)
    paragraph(doc, "scripts/build_reports.py nhận --run-dir và tùy chọn --final-dir để nhập bảng đã lưu, kiểm tra cờ synthetic, fingerprint và selection hash. Script không train lại và không tự diễn giải một mô hình thắng. Phụ lục có thể đưa vào confusion matrix tổng hợp của mô hình đã chọn khi file tồn tại.")
    doc.add_heading("Nguồn tham khảo", level=2)
    for index, (label, url) in enumerate(SOURCES, 1):
        p = paragraph(doc, f"[{index}] {label}. {url}")
        p.paragraph_format.space_after = Pt(5)
        for run in p.runs:
            run.font.size = Pt(8.5)
    paragraph(doc, "Nguồn nội bộ: PLAN.md, configs/experiments.yaml, src/safeview_ml và docs/data.md. Cấu hình trong mã là căn cứ cho hành vi thực tế. Tài liệu phương pháp cần được cập nhật nếu phạm vi hoặc quy trình thay đổi.")

    if results:
        doc.add_page_break()
        doc.add_heading("Phụ lục kết quả đã lưu", level=1)
        paragraph(doc, f"Run {results['selection']['run_id']}. Nguồn là các file comparison.csv đã có; script không huấn luyện lại.")
        paragraph(doc, f"{selection_description(results['selection'])} Dataset fingerprint: {results['selection']['dataset_fingerprint']}.")
        add_table(doc, ["Mô hình", "Biểu diễn", "Validation Macro F1"], [[FAMILIES[row["family"]], representation_label(row), f"{float(row['validation_macro_f1']):.4f}"] for row in results["validation"]], [6.5, 4, 6.5])
        if "test" in results:
            add_table(doc, ["Mô hình", "Test Macro F1", "Accuracy"], [[FAMILIES[row["family"]], f"{float(row['macro_f1']):.4f}", f"{float(row['accuracy']):.4f}"] for row in results["test"]], [7, 5, 5])
            figure = Path(results["final_dir"]) / results["selection"]["selected_family"] / "test_confusion_normalized.png"
            if figure.exists():
                doc.add_picture(str(figure), width=Cm(12))
                paragraph(doc, "Confusion matrix test chuẩn hóa của mô hình triển khai đã khóa trước test.")
        else:
            paragraph(doc, "Chưa chỉ định thư mục final đã hoàn tất, nên phụ lục không có điểm test.")
        paragraph(doc, "Cần bổ sung diễn giải, phân tích lỗi, độ bất định và hạn chế trước khi chuyển bản nháp thành báo cáo nộp.")
    doc.save(path)


def slides_content(results):
    content = [
        {"title": "Phân loại bình luận\ntiếng Việt cho SafeView", "cover": True,
         "body": ["BẢN NHÁP PHƯƠNG PHÁP", "LinearSVC, Logistic Regression, PhoBERT và BamiBERT", "Chưa có kết quả thực nghiệm ViHSD" if not results else "Bảng kết quả đã lưu ở phần phụ lục"], "notes": "Nguồn nội bộ: PLAN.md. Đây là bản nháp phương pháp, chưa phải bài nộp hoàn chỉnh."},
        {"title": "Phạm vi đồ án", "intro": "Đầu vào là một bình luận tiếng Việt. Đầu ra gồm CLEAN, OFFENSIVE và HATE.",
         "table": [["Yêu cầu", "Đầu ra trong workflow"], ["EDA và tiền xử lý", "Audit split, bảng và biểu đồ tổng hợp"], ["Bốn mô hình chính", "SVM, LR, PhoBERT và BamiBERT"], ["Đánh giá", "Macro-F1, metrics từng lớp, confusion matrix"], ["Phân tích và ứng dụng", "Learning curves, lỗi, latency và demo SafeView"]], "widths": [360, 760],
         "bottom": "ComplementNB là baseline phụ trong cấu hình legacy nếu cần ba thuật toán truyền thống cho rubric.", "notes": "Nguồn nội bộ: PLAN.md và configs/experiments.yaml. BamiBERT là đích triển khai khai báo trước, không mặc định là họ đạt validation tốt nhất."},
        {"title": "ViHSD từ GitHub vào runtime", "intro": "Đọc ZIP từ GitHub tác giả vào RAM; cache runtime là tùy chọn, không cần thư mục corpus trong repo.",
         "table": [["Mã nhãn", "Tên nhãn"], ["0", "CLEAN"], ["1", "OFFENSIVE"], ["2", "HATE"]], "widths": [330, 790],
         "bottom": "Ghim commit SHA; giữ train/dev/test và đổi tên dev thành validation. Chưa có điểm số thực nghiệm.", "notes": "Nguồn: https://github.com/sonlam1102/vihsd; src/safeview_ml/remote_data.py. Manifest lưu URL, commit, archive/file hashes và fingerprint; không tự thay dataset khi nguồn lỗi. Hugging Face có xác thực là đường tải tùy chọn."},
        {"title": "Tiền xử lý và đặc trưng", "intro": "NFC và khoảng trắng. Giữ dấu, emoji, phủ định và chữ hoa theo cấu hình mặc định.",
         "table": [["Họ mô hình", "Tiền xử lý / biểu diễn"], ["SVM và LR", "TF-IDF word / char / combined"], ["PhoBERT", "PyVi rồi tokenizer pretrained"], ["BamiBERT", "Văn bản trực tiếp rồi tokenizer pretrained"]], "widths": [390, 730],
         "bottom": "PyVi khác VnCoreNLP trong tiền huấn luyện PhoBERT. Lưu tokenizer, preprocessing và revision cùng model.", "notes": "Nguồn: configs/experiments.yaml, transformers_backend.py; https://huggingface.co/vinai/phobert-base; https://huggingface.co/Qualcomm-AI-Research/BamiBERT. TF-IDF chỉ fit train; Transformer fine-tune từ checkpoint gốc, không dùng model đã train ViHSD không rõ split."},
        {"title": "Bốn mô hình và ngân sách tìm kiếm", "intro": "Baseline tìm đặc trưng rồi tune classifier. Transformer chọn checkpoint bằng validation Macro-F1.",
         "table": [["Mô hình", "Thiết lập khởi đầu"], ["LinearSVC", "C: 0,1 / 1 / 5; weight: None / balanced"], ["Logistic Regression", "C: 0,1 / 1 / 5; weight: None / balanced"], ["PhoBERT", "LR 2e-5; 3 epoch; max_length 128"], ["BamiBERT", "LR 2e-5; 3 epoch; max_length 128"]], "widths": [390, 730],
         "bottom": "Transformer: batch 8, accumulation 2, FP16 khi có CUDA. Chưa đo thời gian hay xác định mô hình thắng.", "notes": "Nguồn nội bộ: configs/experiments.yaml và training.py. Baseline min_df 2/3, TF-IDF tối đa 50000 đặc trưng mỗi nhánh. Ghi thời gian, RAM, kích thước artifact. Số đặc trưng TF-IDF không tương đương số tham số Transformer."},
        {"title": "Validation và đích triển khai", "body": ["1  Fit baseline hoặc fine-tune Transformer trên train", "2  Calibration SVM qua fold train; kiểm tra xác suất mọi họ", "3  Ghi riêng mô hình đạt Macro-F1 validation cao nhất", "4  Khóa BamiBERT và policy trước khi xem điểm test"],
         "bottom": "BamiBERT là đích đã khai báo, không mặc định là họ tốt nhất. Test không chọn model hoặc ngưỡng.", "notes": "Nguồn: configs/experiments.yaml deployment_family; selection.json validation_best_family/selected_family. Baseline có learning curve theo cỡ train, Transformer có loss/metrics theo epoch; báo rõ sự khác biệt."},
        {"title": "Nhãn ba lớp và quyết định ẩn", "formula": "p_harm = P(OFFENSIVE) + P(HATE)",
         "body": ["Nhãn ba lớp lấy argmax của ba xác suất", "Ẩn khi p_harm ≥ threshold", "Chọn ngưỡng tối đa hóa F1 harmful trên validation", "Nếu bằng nhau, ưu tiên FPR CLEAN thấp rồi ngưỡng cao"],
         "bottom": "decision_policy.json lưu threshold, định nghĩa score, tiêu chí chọn và revision. Không thêm điều kiện argmax harmful.", "notes": "Nguồn: https://scikit-learn.org/stable/modules/classification_threshold.html và src/safeview_ml/policy.py. Đây là tiêu chí demo mặc định, chưa phải tối ưu mọi chi phí sản phẩm."},
        {"title": "Tích hợp BamiBERT vào SafeView", "body": ["Model Hub lưu artifact; Space chạy model/tokenizer đã khóa", "Gradio classify trả xác suất CLEAN, OFFENSIVE và HATE", "Provider áp dụng p_harm và policy đã khóa", "Đo riêng model thuần, lexicon, routing và độ trễ API"],
         "bottom": "API lỗi phải báo trạng thái. Token giữ phía server. Chưa xác nhận triển khai public hoặc cải thiện so với PhoBERT.", "notes": "Nguồn nội bộ: PLAN.md mục 5; scripts/templates/hf_space/app.py; docs/safeview-integration.md. Tích hợp thực tế cần artifact đã khóa và endpoint đang vận hành."},
        {"title": "Người dùng chạy thực nghiệm", "body": ["Nạp GitHub trên Colab và lưu manifest/commit nguồn", "Chạy EDA, bốn mô hình và xem khoảng 100 lỗi validation", "Khóa pipeline rồi đánh giá test và bất định", "Hoàn thiện Word, notebook có output, slide và demo"],
         "bottom": "Mẫu tổng hợp chỉ kiểm tra phần mềm. Chưa có kết luận về chất lượng bốn mô hình hoặc khả năng chạy API.", "notes": "Nguồn nội bộ: PLAN.md và notebook. Colab Free không bảo đảm GPU/thời lượng; lưu checkpoint và resume cùng run khi cần. Bản nháp chưa phải bài nộp có kết quả thực nghiệm."},
    ]
    if results:
        content[2]["bottom"] = "Giữ train, validation/dev và test gốc. Fingerprint và run đã dùng được ghi cùng bảng ở phụ lục."
        content[4]["bottom"] = "Bảng phụ lục chỉ dùng số liệu đã lưu. Cấu hình thực tế của từng họ nằm trong run/config.json."
        content[8]["title"] = "Hoàn thiện phân tích và demo"
        content[8]["body"] = ["Rà manifest và EDA tương ứng với run đã lưu", "Phân tích khoảng 100 lỗi validation và learning curves", "Bổ sung bất định và hạn chế của kết quả đã khóa", "Hoàn thiện diễn giải Word, notebook, slide và demo"]
        content[8]["bottom"] = "Bảng đã lưu cung cấp số liệu. Nhóm vẫn cần diễn giải và kiểm tra triển khai trước khi nộp."
        content.append({"title": "Validation từ run đã lưu", "intro": f"Run {results['selection']['run_id']}",
                        "table": [["Mô hình", "Biểu diễn", "Macro-F1"]] + [[FAMILIES[row["family"]], representation_label(row), f"{float(row['validation_macro_f1']):.4f}"] for row in results["validation"]], "widths": [490, 340, 290],
                        "bottom": selection_description(results["selection"]), "notes": f"Nguồn local: {results['run_dir']}/comparison.csv. Fingerprint: {results['selection']['dataset_fingerprint']}"})
        if "test" in results:
            content.append({"title": "Test của các pipeline đã khóa", "intro": "Giữ đích triển khai đã khai báo và khóa trước test; báo riêng họ tốt nhất validation.",
                            "table": [["Mô hình", "Macro-F1", "Accuracy"]] + [[FAMILIES[row["family"]], f"{float(row['macro_f1']):.4f}", f"{float(row['accuracy']):.4f}"] for row in results["test"]], "widths": [570, 275, 275],
                            "bottom": "Điểm test dùng để báo cáo. Không dùng bảng này để thay mô hình hay threshold đã chọn.", "notes": f"Nguồn local: {results['final_dir']}/comparison.csv. Đã kiểm tra selection SHA và fingerprint."})
    return content


JS_BUILDER = r'''
import fs from 'node:fs/promises';
import path from 'node:path';
import { pathToFileURL } from 'node:url';
import { Presentation, PresentationFile } from '@oai/artifact-tool';
const config = JSON.parse(await fs.readFile(process.argv[2], 'utf8'));
const { resolvePresentationFont, finalizePresentation } = await import(pathToFileURL(path.join(config.skill, 'container_tools/artifact_tool_utils.mjs')).href);
const font = resolvePresentationFont({fontFamily: 'Arial'});
const deck = Presentation.create({slideSize: {width:1280,height:720}});
const ink='#17334B', muted='#455D6F', background='#FBFCFD';
function text(slide, value, left, top, width, height, size=28, bold=false, color=ink) {
  const shape = slide.shapes.add({geometry:'textbox',position:{left,top,width,height},fill:'none',line:{fill:'none',width:0}});
  shape.text = value;
  shape.text.style = {typeface:font,fontSize:size,bold,color,autoFit:'none'};
  return shape;
}
const tables=[];
for (const [index, item] of config.slides.entries()) {
  const slide=deck.slides.add(); slide.background.fill=background;
  if(item.cover) {
    text(slide,item.title,72,105,1136,205,58,true);
    text(slide,item.body[0],76,345,1125,48,25,true,muted);
    text(slide,item.body[1],76,407,1110,85,31,false,ink);
    text(slide,item.body[2],76,553,1110,62,28,true,muted);
  } else {
    text(slide,item.title,64,42,1152,80,43,true);
    if(item.intro) text(slide,item.intro,68,133,1140,90,28,false,muted);
    if(item.table) {
      const table=slide.tables.add({rows:item.table.length,columns:item.table[0].length,left:72,top:243,width:1120,height:item.table.length>4?300:265,columnWidths:item.widths,values:item.table});
      table.borders.assign({style:'solid',fill:'#D4DDE4',width:1});
      table.cells.block({row:0,column:0,rowCount:item.table.length,columnCount:item.table[0].length}).assign({textStyle:{typeface:font,fontSize:24,color:ink},margins:{left:16,right:16,top:12,bottom:12},anchor:'center'});
      for(let row=0;row<item.table.length;row++) for(let col=0;col<item.table[0].length;col++) {
        const cell=table.getCell(row,col);cell.fill=row===0?'#DCE6EE':row%2?'#FFFFFF':'#F0F4F7';
        if(row===0)cell.text.style={typeface:font,fontSize:24,bold:true,color:ink};
      }
      tables.push(index+1);
    }
    if(item.formula) text(slide,item.formula,74,143,1120,80,39,true);
    if(item.body) {
      let start=item.formula?257:173;
      const gap=item.formula?72:92;
      for(const [row,line] of item.body.entries())text(slide,line,76,start+gap*row,1128,72,30,false,ink);
    }
    if(item.bottom) text(slide,item.bottom,74,588,1127,88,24,false,muted);
  }
  slide.speakerNotes.textFrame.setText(item.notes);
}
await fs.mkdir(config.build,{recursive:true});
const candidate=path.join(config.build,'candidate.pptx');
await(await PresentationFile.exportPptx(deck)).save(candidate);
const finalPath=path.join(config.build,'final','trinh-bay-phuong-phap.pptx');
await fs.mkdir(path.dirname(finalPath),{recursive:true});
await finalizePresentation({workspaceDir:config.build,candidatePath:candidate,finalPath,
 pythonExecutable:config.python,
 integrityValidatorPath:path.join(config.skill,'container_tools/inspect_presentation_package_integrity.py'),
 layoutValidatorPath:path.join(config.skill,'container_tools/inspect_presentation_layout_geometry.py'),
 layoutArgs:['--expected-slide-size-emu','12192000,6858000','--validate-bullet-geometry','--validate-heading-fit',...tables.flatMap(n=>['--require-native-table-slide',String(n)])],
 explicitTotalSlideCount:config.slides.length,requiredNativeTableOwnerSlides:tables,requiredNativeChartOwnerSlides:[],
 fontPolicy:{basis:'design',families:[font]},verifyArtifactToolImport:true,
 receiptPath:path.join(config.build,'validation.json')});
console.log(JSON.stringify({finalPath,slideCount:config.slides.length,font}));
'''


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=ROOT / "reports")
    parser.add_argument("--build-dir", type=Path, help="Private persistent render/QA directory; default temporary directory")
    parser.add_argument("--run-dir", type=Path)
    parser.add_argument("--final-dir", type=Path)
    parser.add_argument("--overwrite", action="store_true")
    parser.add_argument("--skip-render", action="store_true", help="Development only; outputs still need visual QA before delivery")
    parser.add_argument("--node", type=Path, default=Path(os.environ.get("CODEX_ARTIFACT_NODE", DEPENDENCIES / "node/bin/node")))
    parser.add_argument("--node-modules", type=Path, default=Path(os.environ.get("CODEX_ARTIFACT_NODE_MODULES", DEPENDENCIES / "node/node_modules")))
    parser.add_argument("--documents-skill", type=Path, default=Path(os.environ.get("CODEX_DOCUMENTS_SKILL", SKILLS / "documents/26.909.12148/skills/documents")))
    parser.add_argument("--presentations-skill", type=Path, default=Path(os.environ.get("CODEX_PRESENTATIONS_SKILL", SKILLS / "presentations/26.909.12148/skills/presentations")))
    args = parser.parse_args()
    results = saved_results(args.run_dir, args.final_dir)
    output = args.output_dir.resolve()
    output.mkdir(parents=True, exist_ok=True)
    docx = output / "bao-cao-phuong-phap.docx"
    pptx = output / "trinh-bay-phuong-phap.pptx"
    if not args.overwrite and (docx.exists() or pptx.exists()):
        parser.error("Outputs already exist. Use a new --output-dir or explicitly pass --overwrite.")
    if not args.node.exists() or not args.node_modules.exists():
        parser.error("Bundled artifact runtime is unavailable. Supply --node and --node-modules.")
    build = (args.build_dir or Path(tempfile.mkdtemp(prefix="cs114-reports-"))).resolve()
    build.mkdir(parents=True, exist_ok=True)
    # Use a fresh child so finalization never overwrites a previous final file.
    session = Path(tempfile.mkdtemp(prefix="build-", dir=build))
    (session / "node_modules").symlink_to(args.node_modules.resolve(), target_is_directory=True)
    candidate_docx = session / docx.name
    build_docx(candidate_docx, results)
    config = {"skill": str(args.presentations_skill.resolve()), "build": str(session), "python": sys.executable, "slides": slides_content(results)}
    (session / "slides.json").write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")
    (session / "build.mjs").write_text(JS_BUILDER, encoding="utf-8")
    runtime_env = {**os.environ, "RUNTIME_NODE_MODULES": str(args.node_modules.resolve()), "RUNTIME_NODE": str(args.node.resolve()), "RUNTIME_PYTHON": sys.executable}
    runtime_env["PATH"] = os.pathsep.join([str(DEPENDENCIES / "bin/override"), str(DEPENDENCIES / "bin/fallback"), str(args.node.parent), runtime_env.get("PATH", "")])
    subprocess.run([str(args.node), str(session / "build.mjs"), str(session / "slides.json")], check=True, cwd=session, env=runtime_env)
    if not args.skip_render:
        subprocess.run([sys.executable, str(args.documents_skill / "render_docx.py"), str(candidate_docx), "--output_dir", str(session / "docx-preview")], check=True, cwd=ROOT, env=runtime_env)
        subprocess.run([str(args.node), str(args.presentations_skill / "container_tools/render_presentation.mjs"), "--input", str(session / "final" / pptx.name), "--output_dir", str(session / "pptx-preview")], check=True, cwd=ROOT, env=runtime_env)
    shutil.copyfile(candidate_docx, docx)
    shutil.copyfile(session / "final" / pptx.name, pptx)
    print(json.dumps({"docx": str(docx), "pptx": str(pptx), "private_qa_dir": str(session), "status": "method_draft" if results is None else "draft_with_saved_tables", "visual_review_required": True}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
